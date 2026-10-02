"""검증된 DB 데이터로 Word/Excel 보고서를 생성하고, 파일을 다시 열어 검증한다.

모든 수치와 판정은 DB에서 가져오며 LLM이 집계나 판정을 작성하지 않는다.
"""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from .evidence import EvidenceStore
from .models import VERDICT_LABELS_KO, AuditError, Verdict
from .policy import mask_secrets

VERDICT_ORDER = [v.value for v in Verdict]
LABEL = {v.value: VERDICT_LABELS_KO[v] for v in Verdict}
SUMMARY_HEADER = ["판정", "의미", "건수"]
ITEM_HEADER = ["점검항목ID", "대상ID", "항목명", "심각도", "판정"]
RESULT_COLUMNS = ["점검항목ID", "대상ID", "항목명", "심각도", "판정", "판정(한글)", "판정방법",
                  "판정근거", "증거ID", "조치권고", "담당자", "조치상태", "재진단결과",
                  "체크리스트버전", "규칙버전"]
REQUIRED_COLUMNS = ["점검항목ID", "대상ID", "판정", "판정근거", "체크리스트버전"]
DEFAULT_NOTICE = "본 문서는 Local Audit Workstation 기본 템플릿으로 생성되었으며 공식 기관의 양식이 아니다."


class ReportError(AuditError):
    pass


def _m(v) -> str:
    return mask_secrets("" if v is None else str(v))


def collect(storage, run_id, checklists) -> dict:
    run = storage.get_run(run_id)
    if run is None:
        raise ReportError(f"알 수 없는 실행 {run_id}")
    if run["status"] != "COMPLETED":
        raise ReportError(f"실행 {run_id} 의 상태가 {run['status']} 임. COMPLETED 상태의 실행만 보고서로 출력할 수 있음")
    checklist, sha = checklists.load(run["checklist_id"], run["checklist_version"], for_run=False)
    if sha != run["checklist_sha256"]:
        raise ReportError("체크리스트 내용이 실행에 사용된 버전과 다름")
    items = {it["check_id"]: it for it in checklist["items"]}
    created = next(e for e in storage.audit_events(run_id) if e["event"] == "run_created")
    created_details = json.loads(created["details"])
    scope = created_details.get("scope", {})
    results = []
    for r in storage.results(run_id):
        it = items[r["check_id"]]
        results.append({**dict(r), "title": it["title"], "severity": it["severity"],
                        "remediation": it["remediation"], "evidence_ids": json.loads(r["evidence_ids"]),
                        "rule_outcome": json.loads(r["rule_outcome"]) if r["rule_outcome"] else None,
                        "llm_outcome": json.loads(r["llm_outcome"]) if r["llm_outcome"] else None})
    plan = json.loads(run["plan_json"])
    if "tasks" in created_details:  # platform 필터링 이후의 실제 작업 목록
        expected_keys = {tuple(k) for k in created_details["tasks"]}
    else:
        expected_keys = {(c, t) for c in (plan.get("check_ids") or items) for t in plan["targets"]}
    if {(r["check_id"], r["target_id"]) for r in results} != expected_keys:
        raise ReportError(f"실행 결과 {len(results)}건이 계획된 작업 {len(expected_keys)}건과 일치하지 않음")
    counts = Counter(r["verdict"] for r in results)
    calls = storage.model_calls(run_id)
    return {"run": dict(run), "plan": plan, "scope": scope, "checklist": checklist, "items": items,
            "results": results, "counts": {v: counts.get(v, 0) for v in VERDICT_ORDER},
            "evidence": [dict(e) for e in storage.evidence(run_id)],
            "history": [dict(h) for h in storage.history(run_id)],
            "model_calls": {"total": len(calls), "ok": sum(c["ok"] for c in calls),
                            "prompt_tokens": sum(c["prompt_tokens"] or 0 for c in calls),
                            "completion_tokens": sum(c["completion_tokens"] or 0 for c in calls),
                            "token_unreported_calls": sum(c["token_source"] != "reported" for c in calls)}}


# ------------------------------------------------------------------ Word
def _table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    try:
        t.style = "Table Grid"
    except (KeyError, ValueError):
        pass  # 사용자 템플릿에 이 스타일이 없으면 템플릿 기본값을 유지한다
    for cell, h in zip(t.rows[0].cells, header):
        cell.text = h
    for row in rows:
        for cell, v in zip(t.add_row().cells, row):
            cell.text = _m(v)
    return t


def write_docx(m: dict, path: Path, template: Path | None):
    from docx import Document
    doc = Document(str(template)) if template else Document()
    run, cl = m["run"], m["checklist"]
    doc.add_heading("보안 취약점 진단 보고서", 0)
    if not template:
        doc.add_paragraph(DEFAULT_NOTICE)

    doc.add_heading("1. 진단 개요", 1)
    _table(doc, ["항목", "내용"], [
        ["실행 ID", run["run_id"]], ["진단 계획", m["plan"]["plan_id"]],
        ["Engagement", run["engagement_id"]], ["시작", run["created_at"]],
        ["종료", run["updated_at"]], ["상태", run["status"]]])

    doc.add_heading("2. 대상 및 범위", 1)
    doc.add_paragraph(f"승인자: {_m(m['scope'].get('approved_by'))} / 유효기간: "
                      f"{m['scope'].get('valid_until')} / 허용 위험도: {m['scope'].get('allowed_risks')}")
    _table(doc, ["대상ID", "유형", "값"], [[t["target_id"], t["type"], t["value"]]
                                       for t in m["scope"].get("targets", [])
                                       if t["target_id"] in m["plan"]["targets"]])

    doc.add_heading("3. 진단 방법", 1)
    doc.add_paragraph(f"체크리스트: {cl['guide_name']} (버전 {cl['guide_version']}, "
                      f"ID {run['checklist_id']}@{run['checklist_version']}, SHA-256 {run['checklist_sha256']})")
    doc.add_paragraph(f"공식 가이드 여부: {'예' if cl['official'] else '아니오'} — {cl['notice']}")
    doc.add_paragraph("판정은 등록된 도구로 수집한 증거에 결정론적 규칙을 적용하여 내리며, LLM 분석은 보조 "
                      "의견으로만 사용한다. 규칙과 LLM이 불일치하면 수동확인으로 분류한다.")
    mc = m["model_calls"]
    doc.add_paragraph(f"LLM 호출: {mc['total']}회 (성공 {mc['ok']}), API 보고 토큰: 입력 {mc['prompt_tokens']} / "
                      f"출력 {mc['completion_tokens']}, 토큰 미보고 호출 {mc['token_unreported_calls']}회")

    doc.add_heading("4. 종합 결과", 1)
    _table(doc, SUMMARY_HEADER, [[v, LABEL[v], m["counts"][v]] for v in VERDICT_ORDER] +
           [["TOTAL", "합계", sum(m["counts"].values())]])

    doc.add_heading("5. 항목별 판정", 1)
    _table(doc, ITEM_HEADER, [[r["check_id"], r["target_id"], r["title"], r["severity"], r["verdict"]]
                              for r in m["results"]])

    doc.add_heading("6. 증거 및 분석 근거", 1)
    ev_by_id = {e["evidence_id"]: e for e in m["evidence"]}
    for r in m["results"]:
        doc.add_heading(f"{r['check_id']} / {r['target_id']} — {r['verdict']} ({LABEL[r['verdict']]})", 2)
        doc.add_paragraph(f"판정 근거: {_m(r['reason'])}")
        if r["rule_outcome"]:
            doc.add_paragraph(f"규칙 검사: {_m(json.dumps(r['rule_outcome'], ensure_ascii=False))}")
        if r["llm_outcome"]:
            doc.add_paragraph(f"LLM 분석(참고): {_m(r['llm_outcome'].get('verdict'))} — "
                              f"{_m(r['llm_outcome'].get('rationale') or r['llm_outcome'].get('error'))}")
        if r["verdict"] in ("MANUAL_REVIEW", "INCONCLUSIVE"):  # 담당자가 확인할 원문 기준을 함께 제시
            it = m["items"][r["check_id"]]
            criteria = ((("판단기준(원문)", it["pass_criteria"]),) if it["pass_criteria"] == it["fail_criteria"]
                        else (("양호 기준", it["pass_criteria"]), ("취약 기준", it["fail_criteria"])))
            for label, text in (("점검 절차(원문)", it["procedure"]), *criteria,
                                ("상세 설명(원문)", it.get("source_data", {}).get("상세설명", "")),
                                ("주통 가이드 위치", it.get("source_data", {}).get("주통 가이드 위치", ""))):
                if text.strip():
                    doc.add_paragraph(f"{label}: {_m(text)}")
        for eid in r["evidence_ids"]:
            e = ev_by_id[eid]
            doc.add_paragraph(f"증거 {eid}: {e['tool']} ({e['tool_version']}), 수집 {e['collected_at']}, "
                              f"SHA-256 {e['sha256'] or '-'}, "
                              f"{'성공' if e['success'] else '실패: ' + _m(e['error'])}", style="List Bullet"
                              if "List Bullet" in [s.name for s in doc.styles] else None)

    doc.add_heading("7. 조치 권고", 1)
    fails = [r for r in m["results"] if r["verdict"] == "FAIL"]
    _table(doc, ["점검항목ID", "대상ID", "심각도", "조치 권고"],
           [[r["check_id"], r["target_id"], r["severity"], r["remediation"]] for r in fails])

    doc.add_heading("8. 제한 사항", 1)
    open_items = [r for r in m["results"] if r["verdict"] in ("MANUAL_REVIEW", "INCONCLUSIVE")]
    if not cl["official"]:
        doc.add_paragraph(f"사용한 체크리스트는 공식 가이드가 아니다: {cl['notice']}")
    doc.add_paragraph(f"수동확인/진단불가 항목 {len(open_items)}건은 자동 판정되지 않았으며 담당자 확인이 필요하다.")
    _table(doc, ["점검항목ID", "대상ID", "판정", "사유"],
           [[r["check_id"], r["target_id"], r["verdict"], r["reason"]] for r in open_items])

    doc.add_heading("부록 A. 증거 목록", 1)
    _table(doc, ["증거ID", "점검항목ID", "도구", "SHA-256", "성공"],
           [[e["evidence_id"], e["check_id"], e["tool"], e["sha256"] or "-", "Y" if e["success"] else "N"]
            for e in m["evidence"]])
    doc.add_heading("부록 B. 판정 변경 이력", 1)
    _table(doc, ["시각", "점검항목ID", "대상ID", "이전", "변경", "주체", "사유"],
           [[h["changed_at"], h["check_id"], h["target_id"], h["old_verdict"] or "-", h["new_verdict"],
             h["actor"], h["reason"]] for h in m["history"]])
    doc.save(str(path))


# ------------------------------------------------------------------ Excel
def _sheet(wb, name, header):
    """템플릿에 시트/열 순서가 있으면 그대로 사용하고, 없으면 새로 만든다."""
    if name in wb.sheetnames:
        ws = wb[name]
        tmpl_header = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        if any(tmpl_header):
            return ws, [h if h in header else None for h in tmpl_header]
    else:
        ws = wb.create_sheet(name)
    ws.append(header)
    return ws, header


def _write_rows(ws, cols, rows: list[dict]):
    for r in rows:
        ws.append([_m(r.get(c)) if c and isinstance(r.get(c), str) else (r.get(c) if c else None)
                   for c in cols])


def write_xlsx(m: dict, path: Path, template: Path | None):
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = load_workbook(str(template)) if template else Workbook()
    if not template:
        wb.remove(wb.active)

    ws, cols = _sheet(wb, "요약", ["판정", "의미", "건수"])
    _write_rows(ws, cols, [{"판정": v, "의미": LABEL[v], "건수": m["counts"][v]} for v in VERDICT_ORDER] +
                [{"판정": "TOTAL", "의미": "합계", "건수": sum(m["counts"].values())}])

    ws, cols = _sheet(wb, "판정결과", RESULT_COLUMNS)
    _write_rows(ws, cols, [{
        "점검항목ID": r["check_id"], "대상ID": r["target_id"], "항목명": r["title"], "심각도": r["severity"],
        "판정": r["verdict"], "판정(한글)": LABEL[r["verdict"]], "판정방법": r["method"],
        "판정근거": r["reason"], "증거ID": ", ".join(r["evidence_ids"]), "조치권고": r["remediation"],
        "담당자": "", "조치상태": "미조치" if r["verdict"] == "FAIL" else "", "재진단결과": "",
        "체크리스트버전": r["checklist_version"], "규칙버전": r["rule_version"]} for r in m["results"]])
    if "조치상태" in cols:
        from openpyxl.utils import get_column_letter
        col = get_column_letter(cols.index("조치상태") + 1)
        dv = DataValidation(type="list", formula1='"미조치,조치중,조치완료,위험수용"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")

    ws, cols = _sheet(wb, "체크리스트", ["점검항목ID", "항목명", "상태", "적용대상", "점검절차", "양호기준",
                                    "취약기준", "판정방법", "심각도", "조치권고", "출처", "규칙버전"])
    _write_rows(ws, cols, [{
        "점검항목ID": it["check_id"], "항목명": it["title"], "상태": it["item_status"],
        "적용대상": ",".join(it["target_types"]), "점검절차": it["procedure"], "양호기준": it["pass_criteria"],
        "취약기준": it["fail_criteria"], "판정방법": it["verification_method"]["type"],
        "심각도": it["severity"], "조치권고": it["remediation"], "출처": it["source_reference"],
        "규칙버전": it["rule_version"]} for it in m["checklist"]["items"]])

    ws, cols = _sheet(wb, "증거", ["증거ID", "점검항목ID", "대상ID", "도구", "도구버전", "수집시각",
                                 "SHA-256", "크기", "잘림", "성공", "오류"])
    _write_rows(ws, cols, [{
        "증거ID": e["evidence_id"], "점검항목ID": e["check_id"], "대상ID": e["target_id"], "도구": e["tool"],
        "도구버전": e["tool_version"], "수집시각": e["collected_at"], "SHA-256": e["sha256"] or "",
        "크기": e["size"], "잘림": "Y" if e["truncated"] else "N", "성공": "Y" if e["success"] else "N",
        "오류": e["error"] or ""} for e in m["evidence"]])
    wb.save(str(path))


# ------------------------------------------------------------------ 검증
def _docx_tables(path):
    from docx import Document
    doc = Document(str(path))
    return [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]


def verify(m: dict, docx_path: Path, xlsx_path: Path, evidence_store: EvidenceStore) -> list[str]:
    from openpyxl import load_workbook
    problems = []
    expected = {(r["check_id"], r["target_id"]): r["verdict"] for r in m["results"]}

    # Word
    tables = _docx_tables(docx_path)
    summ = next((t for t in tables if t[0] == SUMMARY_HEADER), None)
    items = next((t for t in tables if t[0] == ITEM_HEADER), None)
    if summ is None or items is None:
        problems.append("docx: 요약 표 또는 항목 표가 없음")
        docx_counts = {}
    else:
        docx_counts = {row[0]: int(row[2]) for row in summ[1:] if row[0] in VERDICT_ORDER}
        docx_items = {(row[0], row[1]): row[4] for row in items[1:]}
        if docx_items != expected:
            problems.append("docx: 항목별 판정이 DB 와 다름")

    # Excel
    wb = load_workbook(str(xlsx_path), read_only=True)
    rows = list(wb["판정결과"].iter_rows(values_only=True))
    header = list(rows[0])
    idx = {h: i for i, h in enumerate(header) if h}
    for col in REQUIRED_COLUMNS:
        if col not in idx:
            problems.append(f"xlsx: 필수 열 '{col}' 이(가) 없음")
    if not problems:
        xl_items = {}
        for row in rows[1:]:
            if not any(row):
                continue
            for col in REQUIRED_COLUMNS:
                if row[idx[col]] in (None, ""):
                    problems.append(f"xlsx: {row[idx['점검항목ID']]} 의 필수 필드 '{col}' 이(가) 비어 있음")
            xl_items[(row[idx["점검항목ID"]], row[idx["대상ID"]])] = row[idx["판정"]]
            if "증거ID" in idx and row[idx["증거ID"]]:
                for eid in str(row[idx["증거ID"]]).split(", "):
                    rec = evidence_store.storage.evidence(evidence_id=eid)
                    problem = "알 수 없는 증거" if rec is None else evidence_store.verify(rec)
                    if problem:
                        problems.append(f"xlsx: 증거 {eid}: {problem}")
        if xl_items != expected:
            problems.append("xlsx: 항목별 판정이 DB 와 다름")
    xl_counts = {r[0]: r[2] for r in wb["요약"].iter_rows(min_row=2, values_only=True) if r[0] in VERDICT_ORDER}
    wb.close()

    if set(m["counts"]) != set(VERDICT_ORDER):
        problems.append("내부 오류: 다섯 가지 판정 상태가 모두 집계되지 않음")
    if docx_counts != m["counts"]:
        problems.append(f"docx 집계 {docx_counts} 가 DB {m['counts']} 와 다름")
    if xl_counts != m["counts"]:
        problems.append(f"xlsx 집계 {xl_counts} 가 DB {m['counts']} 와 다름")
    if sum(m["counts"].values()) != len(expected):
        problems.append("집계 합계가 결과 건수와 다름")
    return problems


def _free(path: Path) -> Path:
    """기존 보고서를 덮어쓰지 않는다."""
    n, p = 1, path
    while p.exists():
        n += 1
        p = path.with_name(f"{path.stem}-{n}{path.suffix}")
    return p


def generate(storage, run_id, checklists, settings, evidence_store) -> dict:
    m = collect(storage, run_id, checklists)
    out = settings.path("reports_dir") / run_id
    out.mkdir(parents=True, exist_ok=True)
    tdir = settings.path("templates_dir")
    wt, xt = tdir / "word" / "report_template.docx", tdir / "excel" / "report_template.xlsx"
    docx_path, xlsx_path = _free(out / f"report_{run_id}.docx"), _free(out / f"report_{run_id}.xlsx")
    write_docx(m, docx_path, wt if wt.exists() else None)
    write_xlsx(m, xlsx_path, xt if xt.exists() else None)
    problems = verify(m, docx_path, xlsx_path, evidence_store)
    result = {"run_id": run_id, "docx": str(docx_path), "xlsx": str(xlsx_path),
              "counts": m["counts"], "problems": problems}
    _free(out / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2),
                                                encoding="utf-8")
    storage.audit(run_id, "report_generated", result)
    return result
