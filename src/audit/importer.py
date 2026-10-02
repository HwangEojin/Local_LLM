"""공식 평가기준 xlsx 를 체크리스트 JSON(draft)으로 변환한다.

원칙: 원문을 요약·수정·보강하지 않는다. 원본에 없는 기준은 만들지 않고, 판단기준을 얻지 못한 항목은
item_status=incomplete 로 표시한다. 자동 판정 수집 도구가 없으므로 모든 항목은 verification_method=manual 이다.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
from pathlib import Path

from . import guide_pdf
from .checklist_engine import validate_content
from .models import ConfigError

# 시트 이름 -> (checklist_id 접미사, 대상 유형). target_types 는 엔진의 라우팅 값이며 원본 내용이 아니다.
SHEETS = {
    "정보보호 관리체계": ("fism", ["host", "path", "url"]),
    "가상화 시스템 관리체계": ("virt-mgmt", ["host", "path", "url"]),
    "클라우드 관리체계": ("cloud-mgmt", ["host", "path", "url"]),
    "서버": ("server", ["host"]),
    "웹서버-WAS": ("was", ["host"]),
    "데이터베이스": ("db", ["host"]),
    "네트워크 인프라": ("net-infra", ["host"]),
    "네트워크 장비": ("net-device", ["host"]),
    "정보보호시스템 장비": ("sec-device", ["host"]),
    "OS 가상화 시스템": ("os-virt", ["host"]),
    "컨테이너 가상화 시스템": ("container", ["host"]),
    "웹_모바일_HTS": ("web-mobile-hts", ["url", "host"]),
}
ID_PREFIX = "fsec-2026-1-"
HEADER_ROW_MARK = "평가항목ID"
SKIP_SHEETS = {"표지"}
# 원본의 오타: 평가대상 열은 OCP_woker, 판단 열은 OCP_worker
SUFFIX_ALIASES = {"OCP_woker": "OCP_worker"}
# 열 종류 -> source_data 키(원본 용어 그대로)
_COMMON = {"평가항목ID", "구분", "통제분야", "통제구분(대)", "통제구분(중)", "평가항목", "위험도", "상세설명",
           "평가기반(전자금융)", "평가기반(주요정보)"}
_ROLE_RE = re.compile(r"^(평가대상|판단기준|판단방법|평가방법|확인 자료(?:\(예시\))?|평가 구분)(?: ?\((.+)\))?$")
_MARK_RE = re.compile(r"(?m)^[ \t]*[\*•o]?[ \t]*(양호|취약|미흡)\b")  # 원본이 취약 대신 미흡이라 쓰기도 함


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "-", s).strip("-") or "x"


def sha256_file(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def split_criteria(text: str) -> tuple[str, str] | None:
    """'* 양호 - ... * 취약 - ...' 서식에서 양호/취약 구간을 원문 그대로 분리한다.

    양호/취약이 각각 정확히 한 구간일 때만 분리한다. 서비스별 구간이 여러 개이거나(예: [S3], [EC2])
    서식이 다르면 None 을 반환하며, 호출 측이 원문 전체를 보존한다(구간을 이어 붙이면 맥락이 사라지므로).
    """
    marks = list(_MARK_RE.finditer(text))
    pass_parts, fail_parts = [], []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        (pass_parts if m.group(1) == "양호" else fail_parts).append(text[m.start():end].strip())
    if len(pass_parts) != 1 or len(fail_parts) != 1:
        return None
    return "\n".join(pass_parts), "\n".join(fail_parts)


class _Sheet:
    """헤더 이름으로 열 역할을 해석한다. 시트마다 열 구성이 다르므로 위치를 가정하지 않는다."""

    def __init__(self, ws):
        self.ws = ws
        header_row = next((r for r in ws.iter_rows(min_row=1, max_row=15)
                           if any(_norm(c.value).startswith(HEADER_ROW_MARK) for c in r)), None)
        if header_row is None:
            raise ConfigError(f"시트 '{ws.title}' 에서 헤더 행('{HEADER_ROW_MARK}')을 찾을 수 없음")
        self.header_row = header_row[0].row
        self.headers = {c.column - 1: _norm(c.value) for c in header_row if _norm(c.value)}
        self.roles: dict[tuple[str, str | None], int] = {}  # (역할, platform) -> 열 인덱스
        self.flag_cols: dict[str, int] = {}  # 접미사 없는 platform 표시 열(예: 스위치)
        self.plain: dict[str, int] = {}  # 공통 열
        last_common = max((i for i, h in self.headers.items() if h.replace(" ", "") in _COMMON), default=0)
        first_role = min((i for i, h in self.headers.items() if _ROLE_RE.match(h) and
                          not h.startswith("평가대상")), default=10**6)
        first_related = min((i for i, h in self.headers.items() if h.startswith("주요 정보통신기반시설")),
                            default=10**6)
        for i, h in sorted(self.headers.items()):
            m = _ROLE_RE.match(h)
            if m:
                suffix = SUFFIX_ALIASES.get(m.group(2), m.group(2))
                key_role = "확인 자료" if m.group(1).startswith("확인 자료") else m.group(1)
                self.roles[(key_role, suffix)] = i
            elif h.replace(" ", "") in _COMMON:
                self.plain[h.replace(" ", "")] = i
            elif last_common < i < min(first_role, first_related) and not h.startswith(("[조항", "전자금융")):
                self.flag_cols[h] = i

    def platforms(self, role: str) -> list[str]:
        return [p for (r, p) in self.roles if r == role and p]

    def has_role(self, role: str) -> bool:
        return any(r == role for (r, _) in self.roles)


def _cell(row, idx):
    return row[idx].value if idx is not None and idx < len(row) else None


def _extras(sheet: _Sheet, row, consumed: set[int]) -> dict[str, str]:
    """사용하지 않은 열(법령 조항, 관련 주통 평가기준 등)을 원문 그대로 보존한다."""
    out = {}
    for i, h in sheet.headers.items():
        if i in consumed or not _text(_cell(row, i)):
            continue
        key = h
        if h.startswith("주요 정보통신기반시설"):
            tag = re.search(r"\[([^\]]+)\]\s*$", h)
            key = "관련 주통 평가기준" + (f" [{tag.group(1)}]" if tag else "")
        out[key] = _text(_cell(row, i))
    return out


def _build_item(sheet: _Sheet, row, rec: dict, check_id: str, platforms: list[str] | None,
                criteria_platform: str | None, target_types: list[str], source_ref: str,
                notes: list[str]) -> dict:
    p = sheet.plain
    crit_col = sheet.roles.get(("판단기준", criteria_platform))
    meth_col = sheet.roles.get(("판단방법", criteria_platform)) or sheet.roles.get(("평가방법", criteria_platform))
    crit = _text(_cell(row, crit_col))
    procedure = _text(_cell(row, meth_col))
    split = split_criteria(crit) if crit else None
    single = bool(crit) and split is None  # 양호/취약 구분 없이 하나의 서술로만 기재된 경우
    if single:
        split = (crit, crit)  # 원문 전체를 그대로 양쪽에 둔다(해석·보강하지 않음)
    status = "complete" if split else "incomplete"
    if status == "incomplete":
        notes.append(f"{check_id}: " + ("판단기준 열 없음" if crit_col is None else "판단기준 비어 있음"))
    data = {k: _text(_cell(row, p[k])) for k in ("구분", "통제분야", "통제구분(대)", "통제구분(중)", "상세설명",
                                                "평가기반(전자금융)", "평가기반(주요정보)") if k in p}
    proof = sheet.roles.get(("확인 자료", criteria_platform))
    if proof is not None and _text(_cell(row, proof)):
        data["확인 자료"] = _text(_cell(row, proof))
    kind = sheet.roles.get(("평가 구분", criteria_platform))
    if kind is not None and _text(_cell(row, kind)):
        data["평가 구분"] = _text(_cell(row, kind))
    if single:
        data["판단기준 서식"] = "양호/취약을 한 쌍으로 분리할 수 없음(단일 서술 또는 서비스별 다중 구간). 원문 전체를 양호/취약 기준 양쪽에 동일하게 수록"
    data["위험도(원문)"] = _text(rec["risk"])
    data.update(rec.get("extra_data", {}))
    data.update(_extras(sheet, row, rec["consumed"]))
    risk = _text(rec["risk"])
    return {
        "check_id": check_id, "title": rec["title"], "item_status": status,
        "source_reference": source_ref,
        "target_types": target_types,
        "applicability": f"대상 platform: {', '.join(platforms)}" if platforms else "",
        "procedure": procedure,
        "pass_criteria": split[0] if split else "", "fail_criteria": split[1] if split else "",
        "evidence_requirements": [], "verification_method": {"type": "manual"},
        "remediation": "", "severity": f"risk-{risk}" if risk in ("1", "2", "3", "4", "5") else "info",
        "rule_version": "source-2026-1",
        **({"platforms": platforms} if platforms else {}),
        "source_data": {k: v for k, v in data.items() if v},
    }


def convert_sheet(ws, filename: str, notes: list[str]) -> list[dict]:
    _, target_types = SHEETS[ws.title]
    sheet = _Sheet(ws)
    web = ws.title == "웹_모바일_HTS"
    per_platform_criteria = any(sheet.roles.get(("판단기준", p)) is not None or
                                sheet.roles.get(("판단방법", p)) is not None
                                for p in sheet.platforms("평가대상") + sheet.platforms("평가 구분") +
                                sheet.platforms("판단기준") + sheet.platforms("판단방법"))
    target_cols = {p: i for (r, p), i in sheet.roles.items() if r == "평가대상" and p}
    items = []
    for row in ws.iter_rows(min_row=sheet.header_row + 1):
        rnum = row[0].row
        title_idx = next((i for i, h in sheet.headers.items() if h.replace(" ", "") == "평가항목"), None)
        title = _text(_cell(row, title_idx))
        if not title:
            continue
        risk_idx = sheet.plain.get("위험도")
        consumed = {i for i in (sheet.plain.get(k) for k in _COMMON) if i is not None}
        consumed |= set(sheet.roles.values()) | set(sheet.flag_cols.values())  # 역할 열은 별도 필드로 옮기므로 제외
        rec = {"title": title, "risk": _cell(row, risk_idx), "consumed": consumed}
        ref = f"{filename} / 시트 '{ws.title}' / 행 {rnum}"

        if web:  # ID 가 채널(WEB/Mobile/HTS)별 접두어 + Sub NUM 으로 나뉘어 있다
            sub = _text(_cell(row, next(i for i, h in sheet.headers.items() if h.startswith("Sub"))))
            for i, h in sorted(sheet.headers.items()):
                m = re.match(r"^평가항목ID \((WEB|Mobile|HTS)\)$", h)
                if m and _text(_cell(row, i)):
                    rec["consumed"] = consumed | {j for j, hh in sheet.headers.items()
                                                  if hh.startswith(("평가항목ID", "Sub"))}
                    items.append(_build_item(sheet, row, rec, f"{_text(_cell(row, i))}{sub}", [m.group(1)],
                                             None, target_types, ref, notes))
            continue

        ident = _text(_cell(row, sheet.plain.get("평가항목ID")))
        if not ident:
            continue
        # platform 해석: 평가대상 열, 평가 구분 열(클라우드), 접미사 없는 표시 열(네트워크 장비)
        flagged = [p for p, i in target_cols.items() if _text(_cell(row, i)).lower() == "o"]
        for (r, p), i in sheet.roles.items():
            if r == "평가 구분" and p and _text(_cell(row, i)) not in ("", "N/A"):
                flagged.append(p)
        flagged += [h for h, i in sheet.flag_cols.items() if _text(_cell(row, i)).lower() == "o"]
        has_platforms = bool(target_cols or sheet.flag_cols or sheet.platforms("평가 구분"))
        if has_platforms and not flagged:
            notes.append(f"{ident}: 적용 platform 표시가 없어 platform 제한 없이 가져옴")
        if per_platform_criteria and flagged:
            # 판단기준 열이 있는 platform 만 펼친다. 열이 없는 표시(웹서버 종류 등)는 적용 대상 정보로 원문 보존.
            crit_platforms = (set(sheet.platforms("판단기준")) | set(sheet.platforms("판단방법")) |
                              set(sheet.platforms("평가방법")))
            expand = [p for p in flagged if p in crit_platforms]
            attr = [p for p in flagged if p not in crit_platforms]
            if attr:
                rec["extra_data"] = {"적용 대상 구분(판단기준 열 없음)": ", ".join(attr)}
            for p in expand or [None]:
                items.append(_build_item(sheet, row, rec, f"{ident}.{_slug(p)}" if p else ident,
                                         [p] if p else (attr or None), p, target_types, ref, notes))
        else:
            items.append(_build_item(sheet, row, rec, ident, flagged or None, None, target_types, ref, notes))
    return items


def _attach_guide_pages(items: list[dict], index: dict, pdf_name: str, notes: list[str]) -> None:
    """항목의 '관련 주통 평가기준' 코드를 PDF 쪽 번호에 연결한다. PDF에 없는 코드(정보보호 정책 A-xx 등)는 건너뛴다."""
    for it in items:
        locs = []
        for key, text in it.get("source_data", {}).items():
            if not key.startswith("관련 주통"):
                continue
            for code, title in guide_pdf.parse_refs(text):
                loc, warn = guide_pdf.locate(code, title, index, pdf_name)
                if loc and loc not in locs:
                    locs.append(loc)
                if warn:
                    notes.append(f"(PDF 연결) {warn}")
        if locs:
            it["source_data"]["주통 가이드 위치"] = "\n".join(locs)


def convert(xlsx: Path, version: str = "2026-1", guide_version: str = "제2026-1호 (개정 2025.12.)",
            pdf: Path | None = None) -> tuple[dict[str, dict], list[str]]:
    """{checklist_id: checklist dict}, 변환 경고 목록을 반환한다."""
    from openpyxl import load_workbook
    xlsx = Path(xlsx)
    wb = load_workbook(str(xlsx), data_only=True)  # 수식 열은 저장된 계산값을 사용
    sources = [{"name": xlsx.name, "sha256": sha256_file(xlsx), "role": "평가기준 원본"}]
    if pdf:
        sources.append({"name": Path(pdf).name, "sha256": sha256_file(pdf), "role": "참고 가이드(관련 주통 평가기준)"})
    notes, out = [], {}
    pdf_index = guide_pdf.index_codes(guide_pdf.extract_pages(Path(pdf))) if pdf else None
    unknown = [n for n in wb.sheetnames if n not in SHEETS and n not in SKIP_SHEETS]
    if unknown:
        raise ConfigError(f"알 수 없는 시트가 있어 변환을 중단함(매핑 필요): {unknown}")
    for name in wb.sheetnames:
        if name in SKIP_SHEETS:
            continue
        cid = ID_PREFIX + SHEETS[name][0]
        items = convert_sheet(wb[name], xlsx.name, notes)
        if pdf_index is not None:
            _attach_guide_pages(items, pdf_index, Path(pdf).name, notes)
        out[cid] = {
            "checklist_id": cid, "guide_name": f"전자금융기반시설 보안 취약점 평가 기준 — {name}",
            "guide_version": guide_version, "official": True,
            "notice": (f"원본 파일 '{xlsx.name}' 의 '{name}' 시트를 자동 변환한 체크리스트이다. 원문은 수정하지 않았다. "
                       "자동 판정용 수집 도구/규칙이 없어 모든 항목은 수동확인(manual)으로 분류되며, "
                       "판단기준을 얻지 못한 항목은 incomplete 로 표시했다."),
            "sources": sources, "items": items,
        }
        errs = validate_content(out[cid])
        if errs:
            raise ConfigError(f"변환 결과가 체크리스트 스키마를 통과하지 못함 ({name}):\n  " + "\n  ".join(errs[:15]))
    return out, notes


def write_checklists(checklists: dict[str, dict], checklist_dir: Path, version: str,
                     source_files: list[Path], replace_draft: bool = False) -> list[Path]:
    """versions/<id>/<version>.json 으로 저장하고 manifest 에 draft 로 추가한다.

    기존 파일은 덮어쓰지 않는다. replace_draft=True 이면 같은 id@version 이 아직 draft 인 경우에만 교체한다
    (활성화되었거나 과거 버전이 된 것은 해시가 고정되어 있으므로 항상 거부).
    """
    checklist_dir = Path(checklist_dir)
    manifest_path = checklist_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {"checklists": []}
    targets = [checklist_dir / "versions" / cid / f"{version}.json" for cid in checklists]
    existing_status = {}
    for cid in checklists:
        entry = next((c for c in manifest["checklists"] if c["checklist_id"] == cid), None)
        ver = next((v for v in (entry or {"versions": []})["versions"] if v["version"] == version), None)
        if ver:
            existing_status[cid] = ver["status"]
    locked = [f"{cid}@{version} ({st})" for cid, st in existing_status.items()
              if not (replace_draft and st == "draft")]
    exists = [str(t) for cid, t in zip(checklists, targets) if t.exists() and cid not in existing_status]
    if locked or exists:
        raise ConfigError("이미 있는 체크리스트는 덮어쓰지 않음 (--replace-draft 는 draft 상태만 교체):\n  " +
                          "\n  ".join(locked + exists))
    src_dir = checklist_dir / "sources"
    src_dir.mkdir(parents=True, exist_ok=True)
    for f in source_files:  # 원본은 복사만 하고, 이미 있으면 내용이 같을 때만 재사용
        dest = src_dir / Path(f).name
        if dest.exists():
            if sha256_file(dest) != sha256_file(f):
                raise ConfigError(f"sources/{dest.name} 이(가) 이미 있으며 내용이 다름")
        else:
            shutil.copy2(f, dest)
    written = []
    for cid, data in checklists.items():
        path = checklist_dir / "versions" / cid / f"{version}.json"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        entry = next((c for c in manifest["checklists"] if c["checklist_id"] == cid), None)
        if entry is None:
            entry = {"checklist_id": cid, "versions": []}
            manifest["checklists"].append(entry)
        entry["versions"] = [v for v in entry["versions"] if v["version"] != version]  # 교체되는 draft 제거
        entry["versions"].append({"version": version, "path": f"versions/{cid}/{version}.json",
                                  "status": "draft"})
        written.append(path)
    tmp = manifest_path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(tmp, manifest_path)
    return written
