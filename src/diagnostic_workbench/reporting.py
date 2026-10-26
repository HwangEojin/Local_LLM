"""보고서: 기존 audit.report / audit.pentest_report 를 수정 없이 재사용한다.

LegacyStorage 가 PostgreSQL 기록을 기존 모듈이 읽는 형식(SQLite Storage 의 조회 메서드)으로 보여 준다.
기존 양식에는 '미실행' 판정이 없어 NOT_RUN 은 진단불가(INCONCLUSIVE)로 싣고 사유에 미실행임을 적으며,
실제 커버리지(NOT_RUN 포함)·작업 상태·차단 사유·도구 실패는 Excel 의 '커버리지' 시트에 따로 남긴다.
"""
from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

from audit.checklist_engine import Checklists
from audit.evidence import EvidenceStore
from audit.local_model import LLMError
from audit.models import AuditError
from audit.policy import mask_secrets

from . import db, explore

_POLISH_SCHEMA = {
    "type": "object", "additionalProperties": False, "required": ["result", "recommendation"],
    "properties": {"result": {"type": "string", "minLength": 1, "maxLength": 1200},
                   "recommendation": {"type": "string", "minLength": 1, "maxLength": 1200}},
}
_POLISH_SYSTEM = (
    "Rewrite the following finding's result(rationale) and recommendation(remediation) for a formal Korean "
    "penetration-test report: clearer, more formal wording only. Do not add new facts, evidence, severity "
    "judgments or claims beyond what is given, and do not drop any technical detail. If it already reads like a "
    "finished report, return it unchanged. Reply with JSON only, in Korean."
)


def _polish_finding_text(llm, title: str, result: str, recommendation: str) -> tuple[str, str]:
    """Reporter 에이전트(멀티에이전트 3단계) — 서술만 다듬는다. CLAUDE.md §3.3: 원문(findings.rationale/
    remediation, DB)은 그대로 두고 이 보고서 출력 문장만 바꾼다 — 다듬기가 실패하거나 엉뚱해도 DB 로
    역전파되지 않으므로 원래 판정·근거는 그대로 안전하다. LLM 이 없거나 실패하면 원문을 그대로 쓴다.
    """
    if llm is None:
        return result, recommendation
    text = f"TITLE: {title}\nRESULT: {result}\nRECOMMENDATION: {recommendation}"
    try:
        out = llm.structured_chat([{"role": "system", "content": _POLISH_SYSTEM}, {"role": "user", "content": text}],
                                  _POLISH_SCHEMA, "report_polish")
    except LLMError:
        return result, recommendation
    return out["result"], out["recommendation"]

_TO_LEGACY = {"NEEDS_REVIEW": "MANUAL_REVIEW", "NOT_RUN": "INCONCLUSIVE"}
_RESULTS = (
    "SELECT rci.id, rci.run_id::text AS run_id, ci.check_id, ra.target_id, rci.verdict, rci.decided_by, rci.reason, "
    "rci.rule_outcome, rci.llm_outcome, rci.updated_at, ci.body, cv.checklist_id, cv.version AS checklist_version, "
    "COALESCE((SELECT array_agg(e.id::text ORDER BY e.collected_at) FROM evidence e "
    "          WHERE e.run_checklist_item_id = rci.id), '{}') AS evidence_ids "
    "FROM run_checklist_items rci JOIN checklist_items ci ON ci.id = rci.checklist_item_id "
    "JOIN checklist_versions cv ON cv.id = ci.checklist_version_id "
    "JOIN run_assets ra ON ra.run_id = rci.run_id AND ra.asset_id = rci.asset_id ")
_EVIDENCE = (
    "SELECT e.id::text AS evidence_id, e.run_id::text AS run_id, COALESCE(ci.check_id, '(탐색)') AS check_id, "
    "ra.target_id, e.requirement_id, e.collected_at, a.kind || ':' || a.value AS target_ref, e.tool, e.tool_version, "
    "e.args, e.rel_path, e.sha256, e.size, e.truncated, e.success, e.error "
    "FROM evidence e JOIN assets a ON a.id = e.asset_id "
    "JOIN run_assets ra ON ra.run_id = e.run_id AND ra.asset_id = e.asset_id "
    "LEFT JOIN run_checklist_items rci ON rci.id = e.run_checklist_item_id "
    "LEFT JOIN checklist_items ci ON ci.id = rci.checklist_item_id ")


def _js(v):
    return json.dumps(v, ensure_ascii=False) if v is not None else None


class LegacyStorage:
    def __init__(self, conn, evidence_dir: Path):
        self.conn, self.evidence_dir = conn, Path(evidence_dir)

    def get_run(self, run_id):
        r = self.conn.execute(
            "SELECT r.*, cv.checklist_id, cv.version AS checklist_version, cv.sha256 AS checklist_sha256 "
            "FROM runs r JOIN checklist_versions cv ON cv.id = r.checklist_version_id WHERE r.id = %s",
            (run_id,)).fetchone()
        if r is None:
            return None
        created = json.loads(self.audit_events(run_id)[0]["details"])
        return {"run_id": str(r["id"]), "plan_json": _js(created["plan"]), "plan_sha256": "",
                "checklist_id": r["checklist_id"], "checklist_version": r["checklist_version"],
                "checklist_sha256": r["checklist_sha256"], "engagement_id": r["engagement_id"],
                "status": r["status"], "created_at": r["created_at"].isoformat(timespec="seconds"),
                "updated_at": r["updated_at"].isoformat(timespec="seconds"), "error": r["stop_reason"]}

    def _result(self, r):
        not_run = r["verdict"] == "NOT_RUN"
        return {"run_id": r["run_id"], "check_id": r["check_id"], "target_id": r["target_id"],
                "verdict": _TO_LEGACY.get(r["verdict"], r["verdict"]),
                "method": "manual_review" if r["decided_by"] == "reviewer" else r["body"]["verification_method"]["type"],
                "reason": "미실행(NOT_RUN): 이 항목은 실행되지 않았음" if not_run else r["reason"],
                "rule_outcome": _js(r["rule_outcome"]), "llm_outcome": _js(r["llm_outcome"]),
                "evidence_ids": _js(r["evidence_ids"]), "checklist_id": r["checklist_id"],
                "checklist_version": r["checklist_version"], "rule_version": r["body"]["rule_version"],
                "decided_at": r["updated_at"].isoformat(timespec="seconds")}

    def results(self, run_id):
        return [self._result(r) for r in self.conn.execute(
            _RESULTS + "WHERE rci.run_id = %s ORDER BY ci.check_id, ra.target_id", (run_id,))]

    def get_result(self, run_id, check_id, target_id):
        r = self.conn.execute(_RESULTS + "WHERE rci.run_id=%s AND ci.check_id=%s AND ra.target_id=%s",
                              (run_id, check_id, target_id)).fetchone()
        return self._result(r) if r else None

    def _evidence(self, e):
        args, rel = e.pop("args"), e.pop("rel_path")
        return {**e, "collected_at": e["collected_at"].isoformat(timespec="seconds"), "args_json": _js(args),
                "raw_path": str(self.evidence_dir / rel) if rel else None,
                "truncated": int(e["truncated"]), "success": int(e["success"])}

    def evidence(self, run_id=None, evidence_id=None):
        if evidence_id:
            e = self.conn.execute(_EVIDENCE + "WHERE e.id::text = %s", (evidence_id,)).fetchone()
            return self._evidence(e) if e else None
        return [self._evidence(e) for e in self.conn.execute(
            _EVIDENCE + "WHERE e.run_id = %s ORDER BY e.collected_at, e.id", (run_id,))]

    def history(self, run_id):
        return [{**h, "old_verdict": _TO_LEGACY.get(h["old_verdict"], h["old_verdict"]),
                 "new_verdict": _TO_LEGACY.get(h["new_verdict"], h["new_verdict"]),
                 "changed_at": h["changed_at"].isoformat(timespec="seconds")} for h in self.conn.execute(
            "SELECT vh.id, rci.run_id::text AS run_id, ci.check_id, ra.target_id, vh.old_verdict, vh.new_verdict, "
            "vh.reason, vh.actor, vh.changed_at FROM verdict_history vh "
            "JOIN run_checklist_items rci ON rci.id = vh.run_checklist_item_id "
            "JOIN checklist_items ci ON ci.id = rci.checklist_item_id "
            "JOIN run_assets ra ON ra.run_id = rci.run_id AND ra.asset_id = rci.asset_id "
            "WHERE rci.run_id = %s ORDER BY vh.id", (run_id,))]

    def audit_events(self, run_id):
        return [{"id": e["id"], "run_id": str(e["run_id"]), "ts": e["ts"].isoformat(timespec="seconds"),
                 "event": e["event"], "details": _js(e["details"])} for e in self.conn.execute(
            "SELECT * FROM audit_events WHERE run_id = %s ORDER BY id", (run_id,))]

    def model_calls(self, run_id):
        return [{**c, "ok": int(c["ok"])} for c in self.conn.execute(
            "SELECT * FROM model_calls WHERE run_id = %s ORDER BY id", (run_id,))]

    def audit(self, run_id, event, details):
        db.audit(self.conn, run_id, event, details)
        self.conn.commit()


def coverage(conn, run_id) -> dict:
    """보고서 상태 표시용 사실: 실제 판정 분포, 작업 상태, 차단·실패 내역."""
    q = lambda sql: conn.execute(sql, (run_id,)).fetchall()  # noqa: E731
    return {
        "verdicts": {r["verdict"]: r["n"] for r in q(
            "SELECT verdict, count(*) AS n FROM run_checklist_items WHERE run_id=%s GROUP BY 1 ORDER BY 1")},
        "backlog": {r["status"]: r["n"] for r in q(
            "SELECT status, count(*) AS n FROM backlog_items WHERE run_id=%s GROUP BY 1 ORDER BY 1")},
        "blocked": q("SELECT kind, reason, policy_class, blocked_reason FROM backlog_items WHERE run_id=%s "
                     "AND status='BLOCKED' ORDER BY created_at"),
        "tool_failures": q("SELECT tool, policy_decision, status, error FROM tool_invocations WHERE run_id=%s "
                           "AND status <> 'SUCCEEDED' ORDER BY started_at"),
        "explorations": q("SELECT state, payload->>'tool' AS tool, payload->>'rationale' AS rationale "
                          "FROM exploration_nodes WHERE run_id=%s AND kind='intent' ORDER BY created_at"),
    }


COVERAGE_HEADER = ["구분", "항목", "값", "비고"]
COVERAGE_TITLE = "커버리지 및 제한 사항"


def _coverage_rows(cov: dict) -> list[list]:
    note = "NOT_RUN 은 판정 결과 표에서 INCONCLUSIVE(진단불가)로 표시됨"
    rows = [["판정(실제)", v, n, note if v == "NOT_RUN" else ""] for v, n in cov["verdicts"].items()]
    rows += [["작업 상태", s, n, ""] for s, n in cov["backlog"].items()]
    rows += [["차단된 작업", b["reason"], b["policy_class"] or "", b["blocked_reason"] or ""] for b in cov["blocked"]]
    rows += [["도구 실패/거부", t["tool"], f"{t['policy_decision']}/{t['status']}", t["error"] or ""]
             for t in cov["tool_failures"]]
    rows += [["탐색(가설, 판정 아님)", e["tool"], e["state"], e["rationale"]] for e in cov["explorations"]]
    return rows


def _open_items(cov: dict) -> int:
    return sum(cov["verdicts"].get(v, 0) for v in ("NOT_RUN", "NEEDS_REVIEW", "INCONCLUSIVE"))


def _add_coverage_sheet(xlsx: Path, cov: dict):
    from openpyxl import load_workbook
    wb = load_workbook(str(xlsx))
    ws = wb.create_sheet("커버리지")
    ws.append(COVERAGE_HEADER)
    for row in _coverage_rows(cov):
        ws.append(row)
    wb.save(str(xlsx))


def _add_coverage_section(docx: Path, cov: dict):
    """기본 Word 보고서 끝에 실제 커버리지 표를 덧붙인다(미실행·검토 필요·도구 실패·차단)."""
    from docx import Document
    doc = Document(str(docx))
    doc.add_paragraph().add_run(COVERAGE_TITLE).bold = True
    n = _open_items(cov)
    doc.add_paragraph(f"미확정 항목(미실행/검토 필요/진단불가) {n}건이 남아 있어 진단이 완전히 끝난 것이 아니다." if n
                      else "미확정 항목 없음.")
    rows = _coverage_rows(cov)
    table = doc.add_table(rows=1 + len(rows), cols=len(COVERAGE_HEADER))
    for r, row in enumerate([COVERAGE_HEADER] + rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = _m_text(value)
    doc.save(str(docx))


def _m_text(v) -> str:
    from audit.policy import mask_secrets
    return mask_secrets("" if v is None else str(v))[:500]


_KIND_BY_PLATFORM = {"WEB": "web"}  # audit.pentest_report 가 지원하는 결과보고서 템플릿 유형


def default_kind(conn, run: dict) -> str | None:
    """실행 대상의 platform 으로 결과보고서 템플릿 유형을 고른다. 지원하는 템플릿이 없으면 None."""
    kinds = {_KIND_BY_PLATFORM.get(r["platform"]) for r in conn.execute(
        "SELECT DISTINCT a.platform FROM run_assets ra JOIN assets a ON a.id = ra.asset_id WHERE ra.run_id = %s",
        (run["id"],))}
    return kinds.pop() if len(kinds) == 1 else None


def _finding_images(conn, evidence_dir: Path, finding_id) -> list[Path]:
    """이 finding 에 연결된 증적 중 이미지(스크린샷) 파일의 실제 경로. 직접 연결된 증적뿐 아니라 그
    증적에서 파생된 이미지(예: Playwright MCP 응답 안의 스크린샷, worker.run_tool 이 derived_from 으로
    따로 저장함)도 포함한다. 디스크에 실제로 있는 파일만 돌려준다(없는 걸 있다고 보고서에 넣지 않는다).
    """
    rows = conn.execute(
        "SELECT e.rel_path FROM finding_evidence fe JOIN evidence e ON e.id = fe.evidence_id "
        "WHERE fe.finding_id=%s AND e.success AND e.media_type IN ('image/png','image/jpeg') "
        "UNION "
        "SELECT d.rel_path FROM finding_evidence fe JOIN evidence d ON d.derived_from = fe.evidence_id "
        "WHERE fe.finding_id=%s AND d.success AND d.media_type IN ('image/png','image/jpeg')",
        (finding_id, finding_id)).fetchall()
    out = []
    for r in rows:
        if r["rel_path"]:
            p = evidence_dir / r["rel_path"]
            if p.is_file():
                out.append(p)
    return out


def _findings_for_pentest_report(conn, evidence_dir: Path, run: dict, meta: dict, llm=None) -> tuple[list[dict], dict[str, int]]:
    """findings 테이블(체크리스트 없는 실행이 자율 탐색으로 쌓은 발견 사항) 을 audit.pentest_report._Builder
    가 기대하는 findings 리스트 모양으로 변환한다. check_id/target_id 는 FAIL 판정과 맞추는 용도가 아니라
    _Builder 가 쓰는 필드 이름일 뿐이므로 체크리스트 없이도 채울 수 있다.

    REJECTED(오탐으로 확인됨)는 본문에서 뺀다 — 기각된 후보를 발견 사항처럼 싣는 것은 CLAUDE.md §3.3 위반이다.
    나머지 상태(status 반환값으로 집계)는 모두 싣고, notes 에 실제 검토 상태를 적어 사람이 아직 보지 않은
    CANDIDATE/NEEDS_REVIEW 를 CONFIRMED 와 구분할 수 있게 한다(§9/§10).
    """
    targets = {t["target_id"]: t for t in meta.get("targets", [])}
    rows = conn.execute(
        "SELECT f.*, ra.target_id FROM findings f "
        "JOIN run_assets ra ON ra.run_id = f.run_id AND ra.asset_id = f.asset_id "
        "WHERE f.run_id=%s AND f.status <> 'REJECTED' ORDER BY f.created_at", (run["id"],)).fetchall()
    status_counts: dict[str, int] = {}
    out = []
    for r in rows:
        status_counts[r["status"]] = status_counts.get(r["status"], 0) + 1
        tinfo = targets.get(r["target_id"]) or {"name": r["target_id"], "url": "", "kind": "WEB"}
        notes = (f"검토 확정(reviewer): {r['reviewed_by']}" if r["status"] == "CONFIRMED" and r["reviewed_by"]
                 else f"검증: {r['verification']} ({r['status']}, 사람 최종 검토 전 — CLAUDE.md §9/§10)")
        images = _finding_images(conn, evidence_dir, r["id"])
        cases = [{"title": "증적 스크린샷", "steps": [
            {"text": f"증적 {i + 1}/{len(images)}", "image": str(p)} for i, p in enumerate(images)]}] if images else []
        result, recommendation = _polish_finding_text(llm, r["title"], mask_secrets(r["rationale"]), r["remediation"] or "")
        out.append({
            "check_id": str(r["id"])[:8], "target_id": r["target_id"], "target_name": tinfo["name"],
            "url": tinfo.get("url", ""), "kind": tinfo.get("kind", "WEB"), "name": r["title"],
            "result": result, "path": tinfo.get("url", "") or r["target_id"],
            "notes": notes, "cases": cases, "recommendation": recommendation, "rec_titles": [],
            "reason": r["rationale"],  # 보고서 서술은 다듬어도 이 필드는 항상 원문(§3.3) — 내부 대조/감사용
        })
    return out, status_counts


def generate_autonomous(cfg, conn, run: dict, meta_path, findings_path=None) -> dict:
    """체크리스트 없는(자율 탐색 전용) 실행의 보고서. audit.pentest_report 의 Word 서식·조립 로직
    (_Builder/fill_*/verify)을 그대로 쓰고, 체크리스트의 collect_findings(FAIL 판정 매칭) 만 findings
    테이블에서 직접 읽는 방식으로 바꿔 넣는다 — 서식은 하나만 유지한다.
    """
    from audit import pentest_report, schemas
    from audit.models import free_path
    if run["status"] != "COMPLETED":
        raise AuditError(f"실행 {run['id']} 의 상태가 {run['status']} 임. COMPLETED 상태의 실행만 보고서로 출력할 수 있음")
    if not meta_path:
        raise AuditError("체크리스트 없는 실행도 --meta <report_meta.json> 이 필요함(고객사·프로젝트·대상)")
    run_id = str(run["id"])
    meta = pentest_report._load_json(meta_path, schemas.REPORT_META, "report meta")
    tpl, rec_paths = pentest_report.find_templates(cfg.settings.path("templates_dir"), "web")
    try:  # Reporter(3단계) — 없거나 실패해도 보고서는 원문으로 계속 나와야 한다(CLAUDE.md §7)
        llm = explore.make_llm(SimpleNamespace(run=run, cfg=cfg, conn=conn), "reporter")
    except LLMError:
        llm = None
    findings, status_counts = _findings_for_pentest_report(conn, cfg.evidence_dir, run, meta, llm)
    b = pentest_report._Builder(tpl, pentest_report.load_recommendations(rec_paths), meta, findings)
    if findings_path:
        b.improvements = json.loads(Path(findings_path).read_text(encoding="utf-8")).get("improvements", [])
    b.fill_overview()
    b.fill_summary()
    b.fill_improvements()
    b.fill_details()
    b.fill_protection()
    b.finish()
    out_dir = cfg.settings.path("reports_dir") / run_id
    out_dir.mkdir(parents=True, exist_ok=True)
    out = free_path(out_dir / f"web_pentest_report_{run_id}.docx")
    b.doc.save(str(out))
    problems, warnings = pentest_report.verify(out, findings, b.unmatched)
    rejected_n = conn.execute(
        "SELECT count(*) AS n FROM findings WHERE run_id=%s AND status='REJECTED'", (run["id"],)).fetchone()["n"]
    result = {"docx": str(out), "template": tpl.name, "findings": len(findings),
              "findings_by_status": status_counts, "rejected_excluded": rejected_n,
              "unmatched_recommendations": sorted(set(b.unmatched)),
              "warnings": ["체크리스트 없는 실행: 발견 사항은 findings 테이블(사람 최종 검토 후 반영)에서 가져옴. "
                          "REJECTED 는 본문에서 제외, CONFIRMED 가 아닌 항목은 notes 란에 검토 상태를 표시함"] + warnings,
              "problems": problems}
    db.audit(conn, run["id"], "pentest_report_generated", result)
    conn.commit()
    return result


def generate(cfg, conn, run: dict, kind=None, meta=None, findings=None) -> dict:
    run_id = str(run["id"])
    if run["checklist_version_id"] is None:
        return generate_autonomous(cfg, conn, run, meta, findings)
    storage = LegacyStorage(conn, cfg.evidence_dir)
    checklists = Checklists(cfg.settings.path("checklist_dir"))
    cov = coverage(conn, run["id"])
    if kind:  # 회사 템플릿 기반 모의해킹 결과보고서(Word)
        from audit import pentest_report
        res = pentest_report.generate(storage, run_id, checklists, cfg.settings, kind, meta, findings)
    else:
        from audit import report
        res = report.generate(storage, run_id, checklists, cfg.settings, EvidenceStore(storage, cfg.evidence_dir))
        _add_coverage_sheet(Path(res["xlsx"]), cov)
        _add_coverage_section(Path(res["docx"]), cov)
        from docx import Document  # 다시 열어 덧붙인 내용이 실제로 들어갔는지 확인한다
        from openpyxl import load_workbook
        sheets = load_workbook(res["xlsx"], read_only=True).sheetnames
        if "커버리지" not in sheets or "판정결과" not in sheets:
            res["problems"].append(f"xlsx: 예상 시트가 없음 {sheets}")
        doc = Document(res["docx"])
        if not any(p.text == COVERAGE_TITLE for p in doc.paragraphs) or                 [c.text for c in doc.tables[-1].rows[0].cells] != COVERAGE_HEADER:
            res["problems"].append("docx: 커버리지 절이 없음")
    return {**res, "coverage": cov}
