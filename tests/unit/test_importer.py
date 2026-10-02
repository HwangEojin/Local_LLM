"""공식 xlsx 변환기 테스트. 실제 원본 대신 같은 구조의 합성 xlsx 를 사용한다."""
import json
import unittest
from pathlib import Path

from openpyxl import Workbook

from audit import report
from audit.checklist_engine import Checklists, validate_content
from audit.importer import convert, split_criteria, write_checklists
from audit.models import ConfigError, PolicyDenied
from tests.helpers import Home

CRIT = "* 양호 - 설정이 안전한 경우\n* 취약 - 설정이 안전하지 않은 경우"


def make_xlsx(path: Path):
    wb = Workbook()
    wb.active.title = "표지"
    ws = wb.create_sheet("서버")
    ws.append([])
    ws.append([None, "서버 보안 취약점 평가 기준"])
    ws.append([])
    ws.append([None, "평가항목ID", "구분", "통제분야", "통제구분(대)", "통제구분(중)", "평가항목", "위험도", "상세설명",
               "평가기반\n(전자금융)", "평가기반\n(주요정보)", "평가대상\n(LINUX)", "평가대상\n(WIN)",
               "평가대상\n(AIX)", "판단기준\n(LINUX)", "판단방법\n(LINUX)", "판단기준\n(WIN)", "판단방법\n(WIN)",
               "주요 정보통신기반시설 취약점 분석 평가기준\n(고시) [UNIX]"])
    ws.append([None, "SRV-001", "기술적 보안", "5. 운영 관리", "5.3", "5.3.1", "SNMP 점검", "3", "상세 설명 원문",
               "o", "o", "o", "o", None, CRIT, "# cat /etc/snmpd.conf", "* 양호 - 사용 안 함\n* 취약 - 사용 중",
               "서비스 확인", "(상) U-59 안전한 SNMP 버전 사용"])
    ws.append([None, "SRV-002", "기술적 보안", "5. 운영 관리", "5.3", "5.3.1", "단일 서술 항목", "5", "상세",
               "o", None, "o", None, None, '설정이 없으면 "취약"으로 판단', "cat x", None, None, None])
    ws.append([None, "SRV-003", "기술적 보안", "5. 운영 관리", "5.3", "5.3.1", "기준 비어 있음", "2", "상세",
               "o", None, "o", None, None, None, None, None, None, None])
    ws2 = wb.create_sheet("클라우드 관리체계")
    ws2.append([])
    ws2.append([])
    ws2.append([])
    ws2.append([None, "평가항목ID", "구분", "통제분야", "통제구분(대)", "평가항목", "위험도", "상세설명",
                "평가기반\n(전자금융)", "평가기반\n(주요정보)", "평가 구분 \n(AWS)", "평가 구분 \n(Azure)",
                "평가방법 (AWS)", "평가방법 (Azure)", "확인 자료 (AWS)", "확인 자료 (Azure)",
                "판단기준 (AWS)", "판단기준 (Azure)"])
    ws2.append([None, "PISM-001", "기술적 보안", "5. 운영 관리", "5.3", "통신 암호화", "5", "상세", "o", "o",
                "스크립트", "N/A", "aws 방법", None, "aws 자료", None, CRIT, None])
    ws3 = wb.create_sheet("웹_모바일_HTS")
    ws3.append([])
    ws3.append([])
    ws3.append([])
    ws3.append([None, "평가항목ID\n(WEB)", "평가항목ID\n(Mobile)", "평가항목ID\n(HTS)", "Sub\nNUM", "구분", "통제분야",
                "통제구분(대)", "통제구분(중)", "평가항목", "위험도", "상세설명", "평가기반\n(전자금융)",
                "평가기반\n(주요정보)"])
    ws3.append([None, "WEB-FIN-", "MOB-FIN-", None, "001", "기술적 보안", "5. 운영 관리", "5.8", "5.8.1", "거래 인증",
                "5", "상세", "o", None])
    wb.save(path)


class ImporterTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import tempfile
        cls.tmp = tempfile.TemporaryDirectory()
        cls.xlsx = Path(cls.tmp.name) / "guide.xlsx"
        make_xlsx(cls.xlsx)
        cls.out, cls.notes = convert(cls.xlsx)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def items(self, cid):
        return {i["check_id"]: i for i in self.out[cid]["items"]}

    def test_split_criteria_verbatim(self):
        p, f = split_criteria(CRIT)
        self.assertEqual((p, f), ("* 양호 - 설정이 안전한 경우", "* 취약 - 설정이 안전하지 않은 경우"))
        self.assertEqual(split_criteria("* 양호 - a\n* 미흡 - b"), ("* 양호 - a", "* 미흡 - b"))
        self.assertIsNone(split_criteria('설정이 없으면 "취약"으로 판단'))

    def test_platform_expansion_and_fidelity(self):
        it = self.items("fsec-2026-1-server")
        self.assertEqual(set(it), {"SRV-001.LINUX", "SRV-001.WIN", "SRV-002.LINUX", "SRV-003.LINUX"})
        linux = it["SRV-001.LINUX"]
        self.assertEqual(linux["platforms"], ["LINUX"])
        self.assertEqual(linux["pass_criteria"], "* 양호 - 설정이 안전한 경우")
        self.assertEqual(linux["procedure"], "# cat /etc/snmpd.conf")  # 원문 그대로
        self.assertEqual(linux["severity"], "risk-3")
        self.assertEqual(linux["verification_method"], {"type": "manual"})
        self.assertEqual(linux["evidence_requirements"], [])
        self.assertIn("행 5", linux["source_reference"])
        self.assertEqual(linux["source_data"]["상세설명"], "상세 설명 원문")
        self.assertIn("U-59", linux["source_data"]["관련 주통 평가기준 [UNIX]"])
        # 플랫폼별 기준이 섞이지 않는다
        self.assertIn("사용 안 함", it["SRV-001.WIN"]["pass_criteria"])
        self.assertNotIn("판단기준 (WIN)", json.dumps(linux["source_data"], ensure_ascii=False))

    def test_missing_or_single_criteria_never_invented(self):
        it = self.items("fsec-2026-1-server")
        self.assertEqual(it["SRV-003.LINUX"]["item_status"], "incomplete")
        self.assertEqual(it["SRV-003.LINUX"]["pass_criteria"], "")
        single = it["SRV-002.LINUX"]
        self.assertEqual(single["pass_criteria"], single["fail_criteria"])
        self.assertEqual(single["pass_criteria"], '설정이 없으면 "취약"으로 판단')
        self.assertIn("단일 서술", single["source_data"]["판단기준 서식"])

    def test_cloud_not_applicable_platform_skipped(self):
        it = self.items("fsec-2026-1-cloud-mgmt")
        self.assertEqual(set(it), {"PISM-001.AWS"})  # Azure 는 N/A
        self.assertEqual(it["PISM-001.AWS"]["source_data"]["확인 자료"], "aws 자료")

    def test_web_channels(self):
        it = self.items("fsec-2026-1-web-mobile-hts")
        self.assertEqual(set(it), {"WEB-FIN-001", "MOB-FIN-001"})
        self.assertEqual(it["MOB-FIN-001"]["platforms"], ["Mobile"])
        self.assertEqual(it["WEB-FIN-001"]["item_status"], "incomplete")  # 원본에 판단기준 열 없음

    def test_all_converted_checklists_validate(self):
        for cid, data in self.out.items():
            self.assertEqual(validate_content(data), [], cid)
            self.assertTrue(data["official"])
            self.assertEqual(len(data["sources"][0]["sha256"]), 64)

    def test_write_never_overwrites_and_keeps_originals(self):
        with __import__("tempfile").TemporaryDirectory() as d:
            cdir = Path(d) / "checklist"
            cdir.mkdir()
            (cdir / "manifest.json").write_text(json.dumps({"checklists": []}), encoding="utf-8")
            write_checklists(self.out, cdir, "2026-1", [self.xlsx])
            man = json.loads((cdir / "manifest.json").read_text(encoding="utf-8"))
            self.assertTrue(all(v["status"] == "draft" for c in man["checklists"] for v in c["versions"]))
            self.assertEqual((cdir / "sources" / "guide.xlsx").read_bytes(), self.xlsx.read_bytes())
            with self.assertRaises(ConfigError):
                write_checklists(self.out, cdir, "2026-1", [self.xlsx])

    def test_unknown_sheet_stops_import(self):
        wb = Workbook()
        wb.active.title = "새로운시트"
        p = Path(self.tmp.name) / "bad.xlsx"
        wb.save(p)
        with self.assertRaises(ConfigError):
            convert(p)


class ImportedChecklistRunTest(unittest.TestCase):
    """변환된 체크리스트로 platform 필터링과 보고서까지 실제로 실행한다."""

    def setUp(self):
        self.h = Home()
        xlsx = self.h.path / "guide.xlsx"
        make_xlsx(xlsx)
        out, _ = convert(xlsx)
        write_checklists(out, self.h.path / "checklist", "test-1", [xlsx])
        cl = Checklists(self.h.path / "checklist")
        cl.activate("fsec-2026-1-server", "test-1")
        self.h.checklists = Checklists(self.h.path / "checklist")
        self.plan = {"plan_id": "p", "checklist_id": "fsec-2026-1-server", "checklist_version": "test-1",
                     "targets": ["sample-app"]}

    def tearDown(self):
        self.h.close()

    def set_platform(self, platform):
        from audit import config
        from audit.policy import Policy
        sc = config.load_scope(self.h.path)
        sc["targets"][0].update(type="host", value="srv01.test")  # 서버 항목의 target_types 는 host
        if platform:
            sc["targets"][0]["platform"] = platform
        self.h.policy = Policy(sc, self.h.path)

    def test_platform_required_and_validated(self):
        self.set_platform(None)
        with self.assertRaises(ConfigError) as cm:  # 서버 체크리스트는 platform 이 필요
            self.h.resolve(self.plan)
        self.assertIn("platform", str(cm.exception))
        self.set_platform("SOLARIS")
        with self.assertRaises(ConfigError):
            self.h.resolve(self.plan)

    def test_run_only_matching_platform_items_and_report(self):
        self.set_platform("LINUX")
        resolved = self.h.resolve(self.plan)
        self.assertEqual({i["check_id"] for i, _ in resolved["tasks"]},
                         {"SRV-001.LINUX", "SRV-002.LINUX", "SRV-003.LINUX"})  # WIN 항목은 제외
        run_id = self.h.orchestrator().start(resolved)
        verdicts = {r["check_id"]: r["verdict"] for r in self.h.storage.results(run_id)}
        self.assertEqual(set(verdicts), {"SRV-001.LINUX", "SRV-002.LINUX", "SRV-003.LINUX"})
        # 자동 판정 수단이 없으므로 양호/취약으로 확정되지 않고 모두 수동확인이어야 한다
        self.assertEqual(set(verdicts.values()), {"MANUAL_REVIEW"})
        self.assertEqual(self.h.storage.evidence(run_id), [])  # 도구 호출 없음
        res = report.generate(self.h.storage, run_id, self.h.checklists, self.h.settings, self.h.evidence)
        self.assertEqual(res["problems"], [])
        self.assertEqual(res["counts"]["MANUAL_REVIEW"], 3)
        from docx import Document
        text = "\n".join(p.text for p in Document(res["docx"]).paragraphs)
        self.assertIn("# cat /etc/snmpd.conf", text)  # 담당자가 확인할 원문 절차가 보고서에 실린다
        self.assertIn("* 양호 - 설정이 안전한 경우", text)


if __name__ == "__main__":
    unittest.main()
