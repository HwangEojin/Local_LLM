import json
import unittest
from pathlib import Path

from audit import report
from tests.helpers import Home
from tests.unit.test_checklist_and_run import EXPECTED


class ReportTest(unittest.TestCase):
    def setUp(self):
        self.h = Home()
        self.run_id = self.h.orchestrator().start(self.h.resolve())

    def tearDown(self):
        self.h.close()

    def gen(self):
        return report.generate(self.h.storage, self.run_id, self.h.checklists, self.h.settings,
                               self.h.evidence)

    def test_generate_and_verify(self):
        res = self.gen()
        self.assertEqual(res["problems"], [])
        self.assertEqual(res["counts"], {"PASS": 1, "FAIL": 2, "MANUAL_REVIEW": 2,
                                         "NOT_APPLICABLE": 1, "INCONCLUSIVE": 0})
        # 독립적으로 다시 열기
        from docx import Document
        from openpyxl import load_workbook
        text = "\n".join(p.text for p in Document(res["docx"]).paragraphs)
        self.assertIn("공식 기관의 양식이 아니다", text)
        self.assertNotIn("changeme123", text)  # 보고서에서 비밀값 마스킹
        wb = load_workbook(res["xlsx"])
        rows = list(wb["판정결과"].iter_rows(values_only=True))
        got = {r[0]: r[4] for r in rows[1:]}
        self.assertEqual(got, EXPECTED)
        self.assertEqual(set(wb.sheetnames), {"요약", "판정결과", "체크리스트", "증거"})

    def test_reports_never_overwritten(self):
        a, b = self.gen(), self.gen()
        self.assertNotEqual(a["docx"], b["docx"])
        self.assertTrue(Path(a["docx"]).exists() and Path(b["docx"]).exists())

    def test_tampered_evidence_detected(self):
        ev = [e for e in self.h.storage.evidence(self.run_id) if e["success"]][0]
        Path(ev["raw_path"]).write_bytes(b"tampered")
        self.assertTrue(any("SHA-256" in p for p in self.gen()["problems"]))

    def test_mismatch_detected(self):
        m = report.collect(self.h.storage, self.run_id, self.h.checklists)
        d = self.h.path / "r.docx"
        x = self.h.path / "r.xlsx"
        report.write_docx(m, d, None)
        m2 = json.loads(json.dumps(m, default=str))
        m2["results"][0]["verdict"] = "PASS" if m2["results"][0]["verdict"] != "PASS" else "FAIL"
        report.write_xlsx(m2, x, None)
        self.assertTrue(any("xlsx: 항목별 판정이 DB 와 다름" in p for p in report.verify(m, d, x, self.h.evidence)))

    def test_user_template_column_order_kept(self):
        from openpyxl import Workbook, load_workbook
        tdir = self.h.settings.path("templates_dir") / "excel"
        tdir.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "판정결과"
        ws.append(["판정", "점검항목ID", "대상ID", "판정근거", "체크리스트버전", "고객사메모"])
        wb.save(tdir / "report_template.xlsx")
        res = self.gen()
        self.assertEqual(res["problems"], [])
        rows = list(load_workbook(res["xlsx"])["판정결과"].iter_rows(values_only=True))
        self.assertEqual(rows[0][:2], ("판정", "점검항목ID"))
        self.assertIsNone(rows[1][5])  # 알 수 없는 템플릿 열은 비워둔다

    def test_incomplete_run_refused(self):
        self.h.storage.set_run_status(self.run_id, "ABORTED")
        with self.assertRaises(report.ReportError):
            self.gen()


if __name__ == "__main__":
    unittest.main()
