"""체크리스트 없이 시작하는 자율 진단(CLAUDE.md §3.2/§5.1/§6)을 확인한다.

- 체크리스트 없이 run create 하려면 --llm(use_llm=True) 이 필요하다(코드/정책 정상 거부).
- 체크리스트 없는 실행은 checklist_version_id 가 NULL 이고, run_checklist_items 없이 goal 노드와
  자산만 등록된다.
- explore.propose 는 체크리스트 증거가 하나도 없어도(known 이 비어 있어도) 체크리스트가 없는 실행이면
  모델을 불러 최초 탐색을 제안하고, 그 제안은 증적 인용 없이도(bootstrap) 받아들여진다.
- 보고서는 새 서식을 만들지 않고 audit.pentest_report 의 Word 서식(_Builder)을 그대로 쓴다. findings
  테이블의 발견 사항을 그 서식이 기대하는 모양으로 변환해 넣을 뿐이다.
"""
import json
import tempfile
import unittest
from pathlib import Path

from audit.models import ConfigError
from diagnostic_workbench import config, db, explore, reporting, runs

ROOT = Path(__file__).resolve().parents[2]
HAVE_WEB_TEMPLATE = any((ROOT / "templates" / "word").glob("웹*결과보고서*.docx"))


class _FakeLLM:
    """explore.propose 가 기대하는 최소 인터페이스. 실제 Ollama 를 부르지 않는다."""
    model = "fake-bootstrap"

    def __init__(self, proposals):
        self._proposals = proposals

    def structured_chat(self, messages, schema, purpose):
        self.last_messages = messages
        return {"proposals": self._proposals}


class CreateRunWithoutChecklist(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            cls.cfg = config.load(None)
            cls.conn = db.connect(cls.cfg.dsn)
            db.require_migrated(cls.conn, cls.cfg.migrations_dir)
        except ConfigError as e:
            raise unittest.SkipTest(f"DB 연결/마이그레이션 불가: {e}") from None
        if "baseline-demo" not in {t["target_id"] for t in __import__("json").loads(
                (cls.cfg.home / "config" / "scope.json").read_text(encoding="utf-8"))["targets"]}:
            raise unittest.SkipTest("scope.json 에 baseline-demo 대상이 없음")

    @classmethod
    def tearDownClass(cls):
        if hasattr(cls, "conn"):
            cls.conn.close()

    def test_requires_llm_without_checklist(self):
        with self.assertRaises(ConfigError):
            runs.create_run(self.conn, self.cfg, ["baseline-demo"], "pytest-autonomous",
                            checklist_id=None, use_llm=False)

    def test_creates_goal_and_assets_without_checklist_items(self):
        run = runs.create_run(self.conn, self.cfg, ["baseline-demo"], "pytest-autonomous",
                              checklist_id=None, use_llm=True)
        self.addCleanup(self.conn.commit)  # create_run 이 이미 commit 했으므로 이후 조회만 함
        self.assertIsNone(run["checklist_version_id"])
        rci_count = self.conn.execute(
            "SELECT count(*) AS n FROM run_checklist_items WHERE run_id=%s", (run["id"],)).fetchone()["n"]
        self.assertEqual(rci_count, 0)
        goal = self.conn.execute(
            "SELECT payload FROM exploration_nodes WHERE run_id=%s AND kind='goal'", (run["id"],)).fetchone()
        self.assertIn("자율 진단", goal["payload"]["text"])

    def test_explore_propose_bootstraps_without_prior_evidence(self):
        run = runs.create_run(self.conn, self.cfg, ["baseline-demo"], "pytest-autonomous",
                              checklist_id=None, use_llm=True)
        ctx = runs.build_context(self.cfg, self.conn, run)
        ctx.llm = _FakeLLM([{
            "rationale": "최초 정찰", "target_id": "baseline-demo", "tool": "mcp:frida/list_processes",
            "args": {}, "risk": "read", "expected_evidence": "프로세스 목록",
            "stop_condition": "목록 확인", "based_on_evidence_ids": [],
        }])
        accepted = explore.propose(ctx)
        self.assertEqual(accepted, 1)
        n = self.conn.execute(
            "SELECT count(*) AS n FROM backlog_items WHERE run_id=%s AND kind='exploration'",
            (run["id"],)).fetchone()["n"]
        self.assertEqual(n, 1)

    @unittest.skipUnless(HAVE_WEB_TEMPLATE, "templates/word 의 웹 결과보고서 템플릿이 필요함")
    def test_report_reuses_pentest_report_format_with_db_findings(self):
        run = runs.create_run(self.conn, self.cfg, ["baseline-demo"], "pytest-autonomous",
                              checklist_id=None, use_llm=True)
        asset_id = self.conn.execute("SELECT asset_id FROM run_assets WHERE run_id=%s", (run["id"],)).fetchone()["asset_id"]
        self.conn.execute(
            "INSERT INTO findings (run_id, asset_id, title, severity, rationale, remediation, verification) "
            "VALUES (%s,%s,'테스트 발견 사항','high','신호 일치로 확정됨','패치 적용','RULE_CONFIRMED')",
            (run["id"], asset_id))
        self.conn.execute("UPDATE runs SET status='COMPLETED', ended_at=now() WHERE id=%s", (run["id"],))
        self.conn.commit()
        run = db.resolve_run(self.conn, str(run["id"]))
        with tempfile.TemporaryDirectory() as tmp:
            meta_path = Path(tmp) / "meta.json"
            meta_path.write_text(json.dumps({
                "customer": "테스트고객사", "project": "테스트 프로젝트",
                "targets": [{"target_id": "baseline-demo", "name": "baseline-demo", "url": "http://localhost:9999",
                            "kind": "WEB"}]}, ensure_ascii=False), encoding="utf-8")
            res = reporting.generate(self.cfg, self.conn, run, meta=meta_path)
        self.assertEqual(res["problems"], [])
        self.assertEqual(res["findings"], 1)
        from docx import Document
        doc = Document(res["docx"])
        full = "\n".join(p.text for p in doc.paragraphs) + "\n" + \
            "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn("테스트 발견 사항", full)


if __name__ == "__main__":
    unittest.main()
